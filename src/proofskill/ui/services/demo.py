"""Deterministic demo services implementing the :class:`AssessmentServices` protocol.

These services exist so the Streamlit UI can run the complete candidate and
recruiter journeys without external API credentials (task P0.12 / FR-15).

Honesty rules baked into this module:

* every result is labelled ``provider_mode="demo"`` — nothing here is a live
  model evaluation, and the UI must show that label;
* repository discovery uses local fixtures and performs **no network requests**;
* answer evaluation is a deterministic keyword heuristic, not an AI judgement;
* errors are returned as explicit :class:`ServiceError` values, never converted
  into fake successes.

The real backend services (Prakarsh's service layer) bind ahead of these
fallbacks through :mod:`proofskill.ui.services.provider` when they exist.
"""

from __future__ import annotations

import io
import re
import uuid
from datetime import datetime, timezone
from typing import Any

from proofskill.ui import config
from proofskill.ui.services import fixtures as fx
from proofskill.ui.services.contracts import (
    OP_RETURN_TYPES,
    AnswerEvaluation,
    AssessmentTask,
    CandidateSummary,
    CommunicationAssessment,
    ConsentRecord,
    DiscoveredRepository,
    DiscoveryResult,
    EvidenceRef,
    EvidenceConfidence,
    ExtractedProfile,
    FollowUp,
    GithubLink,
    JobRequirements,
    ProjectClaim,
    ProofCard,
    Question,
    RepositoryFinding,
    RepositoryScope,
    ResumeExtraction,
    ServiceResult,
    SkillClaim,
    SkillRequirement,
    SkillResult,
)

SUPPORTED_RESUME_EXTENSIONS = ("pdf", "docx", "txt")
_PROVIDER_MODE = "demo"

_PURPOSE_TEXT = {
    "repository_discovery": (
        "List the public repositories on the GitHub profile found in your resume."
    ),
    "repository_analysis": (
        "Read a bounded set of files from the repositories you selected, for assessment only."
    ),
}

_HEADING_RE = re.compile(r"^[A-Z][A-Z0-9 &/()-]{1,40}:?$")
_PROJECT_RE = re.compile(r"^project\s*[:\-]\s*(.*)$", re.IGNORECASE)
_SECTION_RE = re.compile(
    r"^(projects?|personal projects|project work|links?)\s*:?\s*$", re.IGNORECASE
)
_PREFERRED_KEYWORDS = (
    "nice to have",
    "good to have",
    "bonus",
    "optional",
    "preferred",
    "desirable",
)
_COMMUNICATION_SIGNALS = (
    "because",
    "first",
    "then",
    "however",
    "for example",
    "so that",
    "instead",
    "trade-off",
)


def _utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _slug(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug or "untitled"


def _new_id(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:12]}"


def _sentences(text: str) -> list[str]:
    parts = re.split(r"(?<=[.!?])\s+|\n+", text)
    return [p.strip() for p in parts if p and p.strip()]


def _excerpt_for(text: str, needle: str, limit: int = 220) -> str | None:
    for sentence in _sentences(text):
        if re.search(rf"(?<![\w-]){re.escape(needle)}(?![\w-])", sentence, re.IGNORECASE):
            return sentence[:limit]
    return None


def normalize_github_url(raw: str) -> GithubLink | None:
    """Normalise a resume-extracted GitHub URL, rejecting everything else.

    Only ``http(s)://github.com/<owner>[/repo]`` links are accepted. Other
    hosts, embedded credentials and unsupported schemes are ignored — an
    extracted link never triggers a network request on its own.
    """
    if not isinstance(raw, str):
        return None
    candidate = raw.strip().rstrip(").,;\"'")
    lowered = candidate.lower()
    if not lowered.startswith(("http://", "https://")):
        return None
    if "@" in candidate.split("//", 1)[-1].split("/", 1)[0]:
        return None  # embedded credentials
    match = fx.GITHUB_URL_PATTERN.fullmatch(candidate)
    if match is None:
        return None
    owner, repo = match.group(1), match.group(2)
    if repo is None:
        return GithubLink(
            url=f"https://github.com/{owner}",
            link_type="profile",
            account=owner,
        )
    return GithubLink(
        url=f"https://github.com/{owner}/{repo}",
        link_type="repository",
        account=owner,
        repo=repo,
    )


def _find_github_links(text: str) -> tuple[GithubLink, ...]:
    found: dict[str, GithubLink] = {}
    for raw in re.findall(r"https?://[^\s<>\")']+", text):
        link = normalize_github_url(raw)
        if link is not None:
            found.setdefault(link.url, link)
    return tuple(found.values())


def _matched_skills(text: str) -> list[str]:
    matched: list[str] = []
    for skill in fx.SKILL_VOCABULARY:
        if re.search(rf"(?<![\w-]){re.escape(skill)}(?![\w-])", text, re.IGNORECASE):
            matched.append(skill)
    return matched


def _extract_projects(text: str) -> list[ProjectClaim]:
    projects: list[ProjectClaim] = []
    current: str | None = None
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if _SECTION_RE.match(line):
            if current:
                projects.append(_build_project(current, text))
                current = None
            if line.lower().startswith("project"):
                continue
            continue
        match = _PROJECT_RE.match(line)
        if match:
            if current:
                projects.append(_build_project(current, text))
            current = match.group(1).strip()
            continue
        if current is not None:
            if _HEADING_RE.match(line):
                projects.append(_build_project(current, text))
                current = None
            else:
                current = f"{current} {line}"
    if current:
        projects.append(_build_project(current, text))
    return projects


def _build_project(body: str, full_text: str) -> ProjectClaim:
    name, _, description = body.partition(" - ")
    description = description or body
    technologies = tuple(
        skill
        for skill in _matched_skills(f"{name} {description}")
        if skill in fx.TECH_SKILLS
    )
    excerpt = _excerpt_for(full_text, name.split()[0]) if name else None
    return ProjectClaim(
        name=name.strip(),
        description=description.strip(),
        technologies=technologies,
        evidence_excerpt=excerpt,
    )


def _normalize_weights(requirements: list[SkillRequirement]) -> list[SkillRequirement]:
    total = sum(max(r.weight, 0.0) for r in requirements)
    if total <= 0 or not requirements:
        return requirements
    scaled = [
        SkillRequirement(
            skill_name=r.skill_name,
            weight=round(max(r.weight, 0.0) / total * 100, 1),
            required=r.required,
        )
        for r in requirements
    ]
    drift = round(100.0 - sum(r.weight for r in scaled), 1)
    if drift:
        largest = max(range(len(scaled)), key=lambda i: scaled[i].weight)
        scaled[largest] = SkillRequirement(
            skill_name=scaled[largest].skill_name,
            weight=round(scaled[largest].weight + drift, 1),
            required=scaled[largest].required,
        )
    return scaled


class DemoServices:
    """Deterministic, credential-free implementation of the service protocol."""

    mode = "demo"

    def __init__(self, binding_note: str | None = None) -> None:
        self._binding_note = binding_note

    # -- status ----------------------------------------------------------
    def service_status(self) -> dict[str, str]:
        status = {op: "demo" for op in OP_RETURN_TYPES}
        if self._binding_note:
            status["backend"] = self._binding_note
        return status

    # -- resume ----------------------------------------------------------
    def extract_resume_text(
        self, filename: str, content_type: str, data: bytes, max_size_mb: float
    ) -> ServiceResult:
        del content_type  # validated from actual bytes and extension instead
        name = filename or ""
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
        if ext not in SUPPORTED_RESUME_EXTENSIONS:
            allowed = ", ".join(ext.upper() for ext in SUPPORTED_RESUME_EXTENSIONS)
            return ServiceResult.failure(
                "UNSUPPORTED_FILE",
                f"'{name or 'That file'}' is not supported. Upload a file of type: {allowed}.",
                provider_mode=_PROVIDER_MODE,
            )
        if not data:
            return ServiceResult.failure(
                "EXTRACTION_FAILED",
                "The file is empty. Export your resume again and re-upload it.",
                provider_mode=_PROVIDER_MODE,
            )
        limit_bytes = max(max_size_mb, 0.1) * 1024 * 1024
        if len(data) > limit_bytes:
            return ServiceResult.failure(
                "INVALID_INPUT",
                f"The file is {len(data) / (1024 * 1024):.1f} MB; the limit is "
                f"{max_size_mb:g} MB. Compress the file or remove embedded images.",
                provider_mode=_PROVIDER_MODE,
            )

        if ext == "txt":
            text = data.decode("utf-8", errors="replace")
        elif ext == "pdf":
            text = self._extract_pdf(data)
        else:
            text = self._extract_docx(data)

        if text is None:
            return ServiceResult.failure(
                "EXTRACTION_FAILED",
                "We could not read this file. It may be corrupted, password "
                "protected, or not a real "
                f"{ext.upper()} document. Try re-exporting it.",
                provider_mode=_PROVIDER_MODE,
            )
        if len(text.strip()) < config.resume_text_min_length():
            if ext == "pdf":
                return ServiceResult.failure(
                    "EXTRACTION_FAILED",
                    "No readable text was found in this PDF. It is likely a scanned "
                    "or image-only document; text recognition (OCR) is not enabled "
                    "in this build. Upload a text-based PDF, DOCX or TXT file.",
                    provider_mode=_PROVIDER_MODE,
                )
            return ServiceResult.failure(
                "EXTRACTION_FAILED",
                "The document contains almost no readable text. Check that the file "
                "is not blank and try again.",
                provider_mode=_PROVIDER_MODE,
            )
        warnings = ["Text was extracted locally; scanned documents (OCR) are not supported."]
        return ServiceResult.success(
            ResumeExtraction(
                text=text,
                content_type=f"application/{ext}",
                filename=name,
                warnings=tuple(warnings),
            ),
            provider_mode=_PROVIDER_MODE,
        )

    @staticmethod
    def _extract_pdf(data: bytes) -> str | None:
        try:
            try:
                import pymupdf as fitz  # PyMuPDF >= 1.24
            except ModuleNotFoundError:  # pragma: no cover - legacy package alias
                import fitz  # type: ignore[no-redef]
            with fitz.open(stream=data, filetype="pdf") as document:
                return "\n".join(page.get_text() for page in document)
        except Exception:  # noqa: BLE001 - malformed documents must not crash the UI
            return None

    @staticmethod
    def _extract_docx(data: bytes) -> str | None:
        try:
            from docx import Document

            document = Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except Exception:  # noqa: BLE001
            return None

    def extract_candidate_evidence(self, text: str) -> ServiceResult:
        if not isinstance(text, str) or not text.strip():
            return ServiceResult.failure(
                "INVALID_INPUT",
                "No resume text is available to review yet. Upload a resume first.",
                provider_mode=_PROVIDER_MODE,
            )
        skills = _matched_skills(text)
        claims = tuple(
            SkillClaim(
                skill_name=skill,
                status="candidate_claim",
                evidence=(
                    EvidenceRef(
                        source_type="resume",
                        skill_name=skill,
                        excerpt=_excerpt_for(text, skill),
                    ),
                ),
            )
            for skill in skills
        )
        projects = tuple(_extract_projects(text))
        technologies = tuple(sorted({s for s in skills if s in fx.TECH_SKILLS}))
        links = _find_github_links(text)

        warnings: list[str] = [
            "Demo extraction is rule-based and may miss skills written in other words."
        ]
        if not skills:
            warnings.append(
                "No well-known skills were detected. You can continue; skills may "
                "show as 'Not assessed' until an assessment covers them."
            )
        if not projects:
            warnings.append("No 'Project:' entries were recognised in the resume.")
        if not links:
            warnings.append(
                "No GitHub links were found. You can complete an equivalent "
                "alternative assessment instead."
            )
        limitations = (
            "Everything on this screen comes from your resume text: these are "
            "claims, not verified abilities.",
            "Scanned or image-only documents are not processed in this build.",
        )
        return ServiceResult.success(
            ExtractedProfile(
                skills=claims,
                projects=projects,
                technologies=technologies,
                github_links=links,
                warnings=tuple(warnings),
                limitations=limitations,
            ),
            provider_mode=_PROVIDER_MODE,
        )

    # -- job requirements ------------------------------------------------
    def parse_job_requirements(self, job_title: str, job_description: str) -> ServiceResult:
        title = (job_title or "").strip()
        description = (job_description or "").strip()
        if not title and not description:
            return ServiceResult.failure(
                "INVALID_INPUT",
                "Add a job title or a short description so we know what to assess.",
                provider_mode=_PROVIDER_MODE,
            )
        sentences = _sentences(description)
        requirements: list[SkillRequirement] = []
        for skill in fx.SKILL_VOCABULARY:
            occurrences = [
                s for s in sentences
                if re.search(rf"(?<![\w-]){re.escape(skill)}(?![\w-])", s, re.IGNORECASE)
            ]
            if not occurrences:
                continue
            preferred = all(
                any(key in s.lower() for key in _PREFERRED_KEYWORDS) for s in occurrences
            )
            requirements.append(
                SkillRequirement(
                    skill_name=skill,
                    weight=1.0 if preferred else 3.0,
                    required=not preferred,
                )
            )
        warnings: list[str] = []
        if not description:
            warnings.append(
                "The description was empty, so only the job title was considered. "
                "Add skills text for better requirements."
            )
        if not requirements:
            warnings.append(
                "No recognised skills were found in the description. Add explicit "
                "skill names (for example 'Python', 'SQL') so requirements can be "
                "identified — we do not invent them."
            )
        elif not any(r.required for r in requirements):
            warnings.append(
                "All detected skills were marked as preferred; add wording such as "
                "'required' or 'must have' to distinguish mandatory skills."
            )
        requirements = _normalize_weights(requirements)
        return ServiceResult.success(
            JobRequirements(
                job_title=title or "Untitled role",
                requirements=tuple(requirements),
                job_id=f"{_slug(title or description[:40])}-demo",
                description=description,
                warnings=tuple(warnings),
            ),
            provider_mode=_PROVIDER_MODE,
        )

    # -- interview -------------------------------------------------------
    def generate_questions(
        self,
        profile: ExtractedProfile,
        job: JobRequirements,
        max_questions: int,
        route: str,
    ) -> ServiceResult:
        del route  # repository questions are produced by generate_project_tasks
        claimed = {c.skill_name for c in profile.skills}
        ordered: list[str] = [
            r.skill_name
            for r in sorted(
                job.requirements, key=lambda r: (-r.required, -r.weight, r.skill_name)
            )
        ]
        for skill in claimed:
            if skill not in ordered:
                ordered.append(skill)
        if not ordered:
            return ServiceResult.failure(
                "ASSESSMENT_INCOMPLETE",
                "There are no job requirements or resume skills to base questions "
                "on. Add a job description or resume first.",
                provider_mode=_PROVIDER_MODE,
            )
        limit = max(1, int(max_questions))
        skill_evidence = {c.skill_name: c.evidence for c in profile.skills}
        questions: list[Question] = []
        for index, skill in enumerate(ordered[:limit]):
            bank = fx.QUESTION_BANK.get(skill)
            if bank:
                text, q_type, difficulty, concepts = bank[index % len(bank)]
                question_text = text
            else:
                question_text = fx.GENERIC_QUESTION_TEMPLATE.format(skill=skill)
                q_type, difficulty = "technical", "beginner"
                concepts = ("example", "why", "verify", "limit")
            evidence = skill_evidence.get(skill, ())
            excerpt = next((e.excerpt for e in evidence if e.excerpt), None)
            if excerpt:
                context = f"You mentioned {skill} in your resume: “{excerpt}”"
                source = "resume_claim"
            else:
                context = f"This role requires {skill}, which is not clearly claimed on your resume."
                source = "job_requirement"
            questions.append(
                Question(
                    id=f"q-{index + 1:03d}",
                    skill_name=skill,
                    question_text=question_text,
                    question_type=q_type,
                    difficulty=difficulty,
                    weight=1.0,
                    evidence=evidence,
                    context=context,
                    expected_concepts=tuple(concepts),
                    source=source,
                    allow_skip=True,
                )
            )
        return ServiceResult.success(
            tuple(questions),
            provider_mode=_PROVIDER_MODE,
            warnings=("Demo questions come from a fixed local question bank, not a live model.",),
        )

    def evaluate_answer(self, question: Question, answer: str) -> ServiceResult:
        if not isinstance(answer, str) or not answer.strip():
            return ServiceResult.failure(
                "INVALID_INPUT",
                "Type an answer before submitting, or skip this question.",
                provider_mode=_PROVIDER_MODE,
            )
        if question is None or not getattr(question, "question_text", ""):
            return ServiceResult.failure(
                "INVALID_INPUT",
                "This question is no longer available. Continue to the next question.",
                provider_mode=_PROVIDER_MODE,
            )
        low = answer.lower()
        concepts = tuple(question.expected_concepts) or ("example", "why", "verify", "limit")
        matched = [c for c in concepts if c.lower() in low]
        missing = [c for c in concepts if c not in matched]
        score = int(round(len(matched) / len(concepts) * 4))
        score = max(0, min(4, score))

        if not matched:
            explanation = (
                f"The response did not cover the expected concepts for "
                f"{question.skill_name}, so it provided little usable evidence "
                f"under the demo rubric."
            )
        else:
            explanation = (
                f"Covered {len(matched)} of {len(concepts)} expected concepts for "
                f"{question.skill_name} ("
                + ", ".join(matched)
                + ")."
            )
            if missing:
                explanation += " Not covered: " + ", ".join(missing) + "."

        confidence: EvidenceConfidence = "medium" if len(matched) >= 2 else "low"
        signals = sum(1 for token in _COMMUNICATION_SIGNALS if token in low)
        communication = 3.0 if signals else 2.0
        return ServiceResult.success(
            AnswerEvaluation(
                status="assessed",
                skill_name=question.skill_name,
                score=score,
                max_score=4,
                score_scale="0-4",
                explanation=explanation,
                strengths=tuple(matched[:3]),
                gaps=tuple(missing[:3]),
                evidence_references=(f"answer-{question.id}",),
                evidence_confidence=confidence,
                communication_score=communication,
                communication_note=(
                    "Demo clarity/structure heuristic (0-4). It is not a measure of "
                    "technical ability and ignores accent or speaking style."
                ),
                rubric_version="1.0",
                provider_mode=_PROVIDER_MODE,
            ),
            provider_mode=_PROVIDER_MODE,
        )

    def select_follow_up(
        self,
        question: Question,
        evaluation: AnswerEvaluation,
        assessed_skills: list[str],
        budget: int,
    ) -> ServiceResult:
        if budget <= 0:
            return ServiceResult.success(
                FollowUp(action="finish", reason="The question budget is complete."),
                provider_mode=_PROVIDER_MODE,
            )
        weak = (
            evaluation.status == "assessed"
            and evaluation.score is not None
            and evaluation.score <= 2
        )
        if weak and not question.is_follow_up:
            follow_up = Question(
                id=f"{question.id}-fu",
                skill_name=question.skill_name,
                question_text=fx.FOLLOW_UP_TEMPLATE.format(skill=question.skill_name),
                question_type=question.question_type,
                difficulty=question.difficulty,
                weight=question.weight,
                evidence=question.evidence,
                context=f"Follow-up to your previous answer: {evaluation.explanation}",
                expected_concepts=question.expected_concepts,
                is_follow_up=True,
                source=question.source,
                allow_skip=question.allow_skip,
            )
            return ServiceResult.success(
                FollowUp(
                    action="ask_follow_up",
                    skill_name=question.skill_name,
                    reason=(
                        f"The previous answer scored {evaluation.score}/4, so one "
                        "clarifying question was added."
                    ),
                    question=follow_up,
                ),
                provider_mode=_PROVIDER_MODE,
            )
        remaining = [s for s in assessed_skills if s not in {question.skill_name}]
        return ServiceResult.success(
            FollowUp(
                action="assess_next_skill" if remaining else "finish",
                skill_name=question.skill_name,
                reason="The previous answer was sufficiently covered.",
            ),
            provider_mode=_PROVIDER_MODE,
        )

    # -- GitHub ----------------------------------------------------------
    def record_repository_consent(
        self, links: list[GithubLink], granted: bool, scope: str
    ) -> ServiceResult:
        if not links:
            return ServiceResult.failure(
                "INVALID_INPUT",
                "There are no GitHub links to record consent for.",
                provider_mode=_PROVIDER_MODE,
            )
        if scope not in _PURPOSE_TEXT:
            return ServiceResult.failure(
                "INVALID_INPUT",
                "Unknown consent scope.",
                provider_mode=_PROVIDER_MODE,
            )
        return ServiceResult.success(
            ConsentRecord(
                consent_id=_new_id("consent"),
                scope=scope,
                granted=bool(granted),
                purpose=_PURPOSE_TEXT[scope],
                recorded_at=_utc_now(),
            ),
            provider_mode=_PROVIDER_MODE,
        )

    def discover_repositories(
        self, link: GithubLink, consent: ConsentRecord
    ) -> ServiceResult:
        if consent.scope != "repository_discovery":
            return ServiceResult.failure(
                "CONSENT_REQUIRED",
                "Repository discovery needs your recorded consent. Confirm consent "
                "before we look up repositories.",
                provider_mode=_PROVIDER_MODE,
            )
        if not consent.granted:
            return ServiceResult.failure(
                "CONSENT_REQUIRED",
                "You declined repository discovery, so no repositories were listed.",
                provider_mode=_PROVIDER_MODE,
            )
        if link.account in fx.RATE_LIMITED_ACCOUNTS:
            return ServiceResult.failure(
                "INTEGRATION_UNAVAILABLE",
                "GitHub's API rate limit was reached while listing repositories. "
                "Wait a few minutes and try again, or continue with the "
                "alternative assessment.",
                retryable=True,
                provider_mode=_PROVIDER_MODE,
            )

        limitations: list[str] = []
        repositories: tuple[DiscoveredRepository, ...] = ()
        if link.link_type == "repository":
            key = f"{link.account}/{link.repo}"
            repositories = tuple(
                repo
                for repo in fx.DEMO_ACCOUNTS.get(link.account, ())
                if repo.full_name == key
            )
            limitations.append(
                "A direct repository link only offers that repository. The account's "
                "other repositories are not listed or analysed."
            )
            if not repositories:
                limitations.append(
                    "This repository is private, missing, or not readable in this "
                    "build. Private repository access requires a separate "
                    "least-privilege authorization flow that is not available here."
                )
        else:
            repositories = fx.DEMO_ACCOUNTS.get(link.account, ())
            if not repositories:
                limitations.append("No public repositories were found for this account.")
        limitations.append(
            "Repository metadata does not prove that you authored the code."
        )
        return ServiceResult.success(
            DiscoveryResult(
                account=link.account,
                source_url=link.url,
                repositories=repositories,
                selection_required=True,
                consent_id=consent.consent_id,
                limitations=tuple(limitations),
                provider_mode=_PROVIDER_MODE,
            ),
            provider_mode=_PROVIDER_MODE,
        )

    def confirm_repository_scope(
        self, consent: ConsentRecord, repositories: list[DiscoveredRepository]
    ) -> ServiceResult:
        if consent.scope != "repository_analysis" or not consent.granted:
            return ServiceResult.failure(
                "CONSENT_REQUIRED",
                "Analysis consent is required before any repository file is read. "
                "Confirm your selection to continue.",
                provider_mode=_PROVIDER_MODE,
            )
        if not repositories:
            return ServiceResult.failure(
                "INVALID_INPUT",
                "No repository is selected. Select one or more repositories, or "
                "continue with the alternative assessment.",
                provider_mode=_PROVIDER_MODE,
            )
        findings: list[RepositoryFinding] = []
        for repo in repositories:
            summary, files, language = fx.REPO_FINDINGS.get(
                repo.full_name,
                ("Structure read from repository metadata only.", (), None),
            )
            findings.append(
                RepositoryFinding(
                    repository=repo.full_name,
                    summary=summary,
                    file_paths=files,
                    language=language,
                )
            )
        file_count = sum(len(f.file_paths) for f in findings)
        return ServiceResult.success(
            RepositoryScope(
                consent_id=consent.consent_id,
                repositories=tuple(repositories),
                findings=tuple(findings),
                file_count=file_count,
                limitations=(
                    "Repository content does not independently prove that you "
                    "authored every file.",
                    f"Only a bounded set of files (up to {file_count}) from your "
                    "selected repositories is read; nothing is executed.",
                    "No changes are written, committed or pushed to GitHub.",
                ),
                provider_mode=_PROVIDER_MODE,
            ),
            provider_mode=_PROVIDER_MODE,
        )

    # -- assessment tasks -------------------------------------------------
    def generate_project_tasks(
        self, scope: RepositoryScope, job: JobRequirements
    ) -> ServiceResult:
        if not scope.findings:
            return ServiceResult.failure(
                "ASSESSMENT_INCOMPLETE",
                "No selected repository evidence is available. Select a repository "
                "or continue with the alternative assessment.",
                provider_mode=_PROVIDER_MODE,
            )
        primary = scope.findings[0]
        excerpt_file = next(
            (f for f in primary.file_paths if f"{primary.repository}/{f}" in fx.FILE_EXCERPTS),
            primary.file_paths[0] if primary.file_paths else "",
        )
        excerpt = fx.FILE_EXCERPTS.get(f"{primary.repository}/{excerpt_file}")
        skill_cycle = [r.skill_name for r in job.requirements] or ["Python"]
        evidence = (
            EvidenceRef(
                source_type="repository",
                skill_name=skill_cycle[0],
                source_id=f"{primary.repository}:{excerpt_file}" if excerpt_file else primary.repository,
                excerpt=(excerpt or primary.summary)[:160],
            ),
        )
        tasks: list[AssessmentTask] = [
            AssessmentTask(
                id="t-001",
                task_type="code_explanation",
                prompt=(
                    f"Explain what the code in `{excerpt_file}` does, step by step, "
                    "and why the author wrote it that way. Point out one edge case "
                    "it does not handle."
                ),
                skill_name=skill_cycle[0],
                difficulty="intermediate",
                evidence=evidence,
                file_refs=(f"{primary.repository}/{excerpt_file}",) if excerpt_file else (),
                code_excerpt=excerpt,
                response_kind="text",
                context="Grounded in a bounded excerpt from the repository you selected.",
                expected_concepts=("route", "validation", "response", "edge case"),
            ),
            AssessmentTask(
                id="t-002",
                task_type="code_change",
                prompt=(
                    f"Submit a focused patch (unified diff) for `{excerpt_file}` that "
                    "adds a reasonable input limit or validation, plus a one-paragraph "
                    "explanation of your change."
                ),
                skill_name=skill_cycle[min(1, len(skill_cycle) - 1)],
                difficulty="intermediate",
                evidence=evidence,
                file_refs=(f"{primary.repository}/{excerpt_file}",) if excerpt_file else (),
                code_excerpt=excerpt,
                response_kind="patch",
                context="Your patch is reviewed as an assessment response only.",
                expected_concepts=("diff", "guard", "input", "test"),
            ),
            AssessmentTask(
                id="t-003",
                task_type="approach_comparison",
                prompt=(
                    f"For the storage approach used in `{primary.repository}`, compare "
                    "keeping the current design against moving to a different one. "
                    "Which would you choose for 10x the data, and why?"
                ),
                skill_name=skill_cycle[min(2, len(skill_cycle) - 1)],
                difficulty="advanced",
                evidence=evidence,
                file_refs=(primary.repository,),
                response_kind="text",
                context="Discuss trade-offs rather than a single correct answer.",
                expected_concepts=("trade-off", "scale", "complexity", "choose"),
            ),
            AssessmentTask(
                id="t-004",
                task_type="architecture",
                prompt=(
                    f"Describe the architecture of `{primary.repository}`: the main "
                    "components, how data flows through them, and one change you "
                    "would make as it grows."
                ),
                skill_name=skill_cycle[min(3, len(skill_cycle) - 1)],
                difficulty="advanced",
                evidence=evidence,
                file_refs=(primary.repository,),
                response_kind="text",
                context="Based on the files listed for the repository you selected.",
                expected_concepts=("component", "data", "flow", "change"),
            ),
            AssessmentTask(
                id="t-005",
                task_type="testing",
                prompt=(
                    f"How would you test `{primary.repository}` reliably, including "
                    "the parts that touch a database or external services? Name the "
                    "kinds of tests you would write first."
                ),
                skill_name=skill_cycle[0],
                difficulty="intermediate",
                evidence=evidence,
                file_refs=(primary.repository,),
                response_kind="text",
                context="Repository evidence is limited to the selected scope.",
                expected_concepts=("unit", "mock", "edge", "integration"),
            ),
        ]
        return ServiceResult.success(
            tuple(tasks),
            provider_mode=_PROVIDER_MODE,
            warnings=(
                "Only bounded excerpts from selected repositories are shown; full "
                "repository contents are never displayed.",
            ),
        )

    def create_alternative_assessment(
        self, profile: ExtractedProfile, job: JobRequirements
    ) -> ServiceResult:
        skill_cycle = [r.skill_name for r in job.requirements] or [
            c.skill_name for c in profile.skills
        ] or ["Problem solving"]
        project = profile.projects[0] if profile.projects else None
        project_context = (
            f"Reference: your resume lists “{project.name}” — {project.description}"
            if project
            else "Reference: describe any project you have worked on, even a small one."
        )
        tasks: list[AssessmentTask] = [
            AssessmentTask(
                id="t-101",
                task_type="scenario",
                prompt=(
                    "A teammate's change breaks a feature you own, and you cannot "
                    "reproduce the failure locally. Walk through the steps you would "
                    "take, in order, to find and fix it."
                ),
                skill_name=skill_cycle[0],
                difficulty="intermediate",
                response_kind="text",
                context="Scenario-based question — no repository access required.",
                expected_concepts=("reproduce", "isolate", "logs", "verify"),
            ),
            AssessmentTask(
                id="t-102",
                task_type="design",
                prompt=(
                    "Design a small service that lets users save and search notes. "
                    "Describe the data model, the main endpoints or interfaces, and "
                    "one trade-off you would accept for speed."
                ),
                skill_name=skill_cycle[min(1, len(skill_cycle) - 1)],
                difficulty="intermediate",
                response_kind="text",
                context="Practical design task — no repository access required.",
                expected_concepts=("model", "endpoint", "index", "trade-off"),
            ),
            AssessmentTask(
                id="t-103",
                task_type="walkthrough",
                prompt=(
                    "Walk through one project from your experience: what it does, "
                    "exactly what you built, one difficult decision, and one thing "
                    "you would change."
                ),
                skill_name=skill_cycle[min(2, len(skill_cycle) - 1)],
                difficulty="beginner",
                response_kind="text",
                context=project_context,
                expected_concepts=("built", "decision", "change", "result"),
            ),
            AssessmentTask(
                id="t-104",
                task_type="code_change",
                prompt=(
                    "Write a short function (as text) that takes a list of records "
                    "and returns the top N by a numeric field, including a brief note "
                    "on how you would test it."
                ),
                skill_name=skill_cycle[0],
                difficulty="intermediate",
                response_kind="patch",
                context="Structured coding explanation — paste your code as text.",
                expected_concepts=("sort", "parameter", "return", "test"),
            ),
            AssessmentTask(
                id="t-105",
                task_type="implementation",
                prompt=(
                    f"This role needs {skill_cycle[-1]}. Describe how you would "
                    "approach a first task that requires it, how you would check "
                    "your work, and what you would ask if requirements were unclear."
                ),
                skill_name=skill_cycle[-1],
                difficulty="intermediate",
                response_kind="text",
                context="Role-relevant implementation question.",
                expected_concepts=("approach", "check", "ask", "deliver"),
            ),
        ]
        return ServiceResult.success(
            tuple(tasks),
            provider_mode=_PROVIDER_MODE,
            warnings=(
                "Alternative assessment — equivalent scope and rubric, no repository "
                "evidence required.",
            ),
        )

    # -- proof card -------------------------------------------------------
    def generate_proof_card(
        self,
        job: JobRequirements,
        profile: ExtractedProfile,
        route: str,
        answer_evaluations: list[AnswerEvaluation],
        task_evaluations: list[AnswerEvaluation],
        scope: RepositoryScope | None,
        candidate_id: str,
        candidate_reference: str,
    ) -> ServiceResult:
        targets = [(r.skill_name, r.weight) for r in job.requirements]
        if not targets:
            targets = [(c.skill_name, 1.0) for c in profile.skills]
        if not targets:
            return ServiceResult.failure(
                "ASSESSMENT_INCOMPLETE",
                "There is nothing to summarise yet — no job requirements or resume "
                "skills are available.",
                provider_mode=_PROVIDER_MODE,
            )
        evaluations = list(answer_evaluations) + list(task_evaluations)
        by_skill: dict[str, list[AnswerEvaluation]] = {}
        for ev in evaluations:
            if ev is None or ev.status != "assessed" or ev.score is None:
                continue
            by_skill.setdefault(ev.skill_name, []).append(ev)

        weights = {name: weight for name, weight in targets}
        skill_results = _build_skill_results(targets, by_skill)
        performance, score_status, coverage = _overall(skill_results, weights)
        card = _compose_proof_card(
            candidate_id=candidate_id,
            candidate_reference=candidate_reference,
            job_id=job.job_id,
            job_title=job.job_title,
            route=route,
            skill_results=skill_results,
            weights=weights,
            performance=performance,
            score_status=score_status,
            coverage=coverage,
            evaluations=evaluations,
            profile_skills=[c.skill_name for c in profile.skills],
            scope=scope,
            provider_mode=_PROVIDER_MODE,
            provenance={
                "rubric_version": "1.0",
                "route": route,
                "interview_answers": len(answer_evaluations),
                "task_responses": len(task_evaluations),
                "services": self.service_status(),
            },
        )
        return ServiceResult.success(card, provider_mode=_PROVIDER_MODE)

    # -- recruiter ---------------------------------------------------------
    def list_jobs(self) -> ServiceResult:
        return ServiceResult.success(fx.DEMO_JOBS, provider_mode=_PROVIDER_MODE)

    def list_candidates(self, job_id: str) -> ServiceResult:
        cards = fixture_proof_cards()
        summaries = tuple(
            CandidateSummary(
                candidate_id=candidate_id,
                reference_label=fx.DEMO_CANDIDATE_REFERENCES[candidate_id],
                job_id=card.job_id,
                job_title=card.job_title,
                status="completed",
                proof_card=card,
                provider_mode=_PROVIDER_MODE,
            )
            for candidate_id, card in cards.items()
            if card.job_id == job_id
        )
        return ServiceResult.success(summaries, provider_mode=_PROVIDER_MODE)

    def compare_candidates(self, job_id: str, candidate_ids: list[str]) -> ServiceResult:
        job = next((j for j in fx.DEMO_JOBS if j.job_id == job_id), None)
        if job is None:
            return ServiceResult.failure(
                "NOT_FOUND", "That job description could not be found.",
                provider_mode=_PROVIDER_MODE,
            )
        cards = fixture_proof_cards()
        missing = [
            cid for cid in candidate_ids
            if cid not in cards or cards[cid].job_id != job_id
        ]
        if missing:
            return ServiceResult.failure(
                "NOT_FOUND",
                "One or more selected candidates do not have an assessment for this "
                "job, so they cannot be compared on the same criteria.",
                provider_mode=_PROVIDER_MODE,
            )
        if len(candidate_ids) < 2:
            return ServiceResult.failure(
                "INVALID_INPUT",
                "Select at least two candidates to compare.",
                provider_mode=_PROVIDER_MODE,
            )

        from proofskill.ui.services.recruiter import build_comparison

        return ServiceResult.success(
            build_comparison(job=job, cards=cards, candidate_ids=candidate_ids),
            provider_mode=_PROVIDER_MODE,
        )


# --------------------------------------------------------------------------
# Shared aggregation helpers (used by the demo service and its fixtures)
# --------------------------------------------------------------------------


def _build_skill_results(
    targets: list[tuple[str, float]],
    by_skill: dict[str, list[AnswerEvaluation]],
) -> tuple[SkillResult, ...]:
    results: list[SkillResult] = []
    for skill_name, _weight in targets:
        evs = by_skill.get(skill_name, [])
        if not evs:
            results.append(
                SkillResult(
                    skill_name=skill_name,
                    performance_score=None,
                    evidence_coverage=0.0,
                    evidence_confidence="not_assessed",
                    status="not_assessed",
                    explanation=(
                        "Not assessed — no assessment evidence was collected for "
                        "this skill. This is not a zero."
                    ),
                    question_count=0,
                )
            )
            continue
        scores = [e.score / e.max_score * 100 for e in evs if e.max_score]
        performance = round(sum(scores) / len(scores), 1)
        average_rubric = sum(e.score for e in evs if e.score is not None) / len(evs)
        if len(evs) >= 2 and average_rubric >= 3:
            confidence: EvidenceConfidence = "high"
        elif average_rubric >= 2:
            confidence = "medium"
        else:
            confidence = "low"
        evidence = tuple(
            EvidenceRef(
                source_type="assessment",
                skill_name=skill_name,
                source_id=e.evidence_references[0] if e.evidence_references else None,
                excerpt=e.explanation[:160] or None,
            )
            for e in evs
        )
        results.append(
            SkillResult(
                skill_name=skill_name,
                performance_score=performance,
                evidence_coverage=100.0,
                evidence_confidence=confidence,
                status="assessed",
                evidence=evidence,
                explanation=f"Based on {len(evs)} evaluated answer(s) on a 0-4 rubric.",
                question_count=len(evs),
            )
        )
    return tuple(results)


def _overall(
    skill_results: tuple[SkillResult, ...],
    weights: dict[str, float],
) -> tuple[float | None, str, float]:
    if not skill_results:
        return None, "not_assessed", 0.0
    total_weight = sum(max(weights.get(s.skill_name, 1.0), 0.0) for s in skill_results)
    assessed = [s for s in skill_results if s.status == "assessed" and s.performance_score is not None]
    assessed_weight = sum(max(weights.get(s.skill_name, 1.0), 0.0) for s in assessed)
    coverage = round(100.0 * assessed_weight / total_weight, 1) if total_weight else 0.0
    if not assessed:
        return None, "not_assessed", coverage
    performance = round(
        sum(s.performance_score * max(weights.get(s.skill_name, 1.0), 0.0) for s in assessed)
        / assessed_weight,
        1,
    )
    status = "complete" if len(assessed) == len(skill_results) else "provisional"
    return performance, status, coverage


def _compose_proof_card(
    *,
    candidate_id: str,
    candidate_reference: str,
    job_id: str | None,
    job_title: str,
    route: str,
    skill_results: tuple[SkillResult, ...],
    weights: dict[str, float],
    performance: float | None,
    score_status: str,
    coverage: float,
    evaluations: list[AnswerEvaluation],
    profile_skills: list[str],
    scope: RepositoryScope | None,
    provider_mode: str,
    provenance: dict[str, Any],
    communication_score: float | None = None,
) -> ProofCard:
    unassessed = tuple(s.skill_name for s in skill_results if s.status == "not_assessed")
    assessed = {s.skill_name.lower() for s in skill_results if s.status == "assessed"}
    claims_not_assessed = tuple(
        name for name in profile_skills if name.lower() not in assessed
    )
    strengths = tuple(
        s.skill_name
        for s in sorted(
            (s for s in skill_results if (s.performance_score or 0) >= 70),
            key=lambda s: -(s.performance_score or 0),
        )[:4]
    )
    comm_scores = [e.communication_score for e in evaluations if e.communication_score is not None]
    if communication_score is not None:
        communication = CommunicationAssessment(
            status="assessed",
            score=round(float(communication_score), 1),
            scale="0-100",
            explanation=(
                "Synthetic demo clarity/structure score (0-100). Communication never "
                "substitutes for technical performance."
            ),
        )
    elif comm_scores:
        communication = CommunicationAssessment(
            status="assessed",
            score=round(sum(comm_scores) / len(comm_scores) * 25, 1),
            scale="0-100",
            explanation=(
                "Clarity and structure of answers on a 0-4 rubric, normalised to "
                "0-100. Communication never substitutes for technical performance."
            ),
        )
    else:
        communication = CommunicationAssessment(
            status="not_assessed",
            score=None,
            explanation="No answers were evaluated, so communication was not assessed.",
        )

    limitations: list[str] = []
    if route == "repository" and scope is not None:
        limitations.append(
            f"Bounded read-only inspection of {len(scope.repositories)} selected "
            "repository(s); repository ownership does not prove authorship."
        )
    elif route == "repository":
        limitations.append("Repository evidence was requested but no scope was confirmed.")
    else:
        limitations.append(
            "No repository evidence was used — the equivalent alternative "
            "assessment path was completed."
        )
    if unassessed:
        limitations.append(
            f"{len(unassessed)} skill(s) remain unassessed and are shown as "
            "'Not assessed', never as zero."
        )
    if provider_mode == "demo":
        limitations.append(
            "Generated in demo mode with a deterministic fallback — no live model "
            "was used for these results."
        )

    follow_ups: list[str] = [
        f"Run a focused assessment for {name}." for name in unassessed
    ]
    follow_ups.extend(
        f"Add practical evidence for {s.skill_name} (confidence is "
        f"{s.evidence_confidence})."
        for s in skill_results
        if s.status == "assessed" and s.evidence_confidence == "low"
    )
    evidence = tuple(
        ref for s in skill_results for ref in s.evidence if ref.source_type == "assessment"
    )
    return ProofCard(
        candidate_id=candidate_id,
        candidate_reference=candidate_reference,
        job_id=job_id,
        job_title=job_title,
        assessment_route=route,
        performance_score=performance,
        score_status=score_status,
        skill_results=skill_results,
        communication=communication,
        evidence_coverage=coverage,
        score_scale="0-100",
        strengths=strengths,
        limitations=tuple(limitations),
        unassessed_skills=unassessed,
        claims_not_assessed=claims_not_assessed,
        recommended_follow_ups=tuple(follow_ups),
        evidence=evidence,
        rubric_version=str(provenance.get("rubric_version", "1.0")),
        provider_mode=provider_mode,  # type: ignore[arg-type]
        provenance=provenance,
        generated_at=_utc_now(),
        is_demo=provider_mode == "demo",
    )


# --------------------------------------------------------------------------
# Synthetic recruiter fixtures
# --------------------------------------------------------------------------

_FIXTURE_SPECS: dict[str, dict[str, Any]] = {
    "demo-candidate-001": {
        "job_id": "demo-job-python-001",
        "route": "repository",
        "communication": 78.0,
        "skills": {
            "Python": (82.5, "medium"),
            "Problem solving": (70.0, "medium"),
            "Data structures": (62.5, "low"),
            "SQL": (None, "not_assessed"),
            "Testing": (None, "not_assessed"),
        },
    },
    "demo-candidate-002": {
        "job_id": "demo-job-python-001",
        "route": "alternative",
        "communication": 81.0,
        "skills": {
            "Python": (76.0, "medium"),
            "Problem solving": (84.0, "high"),
            "Data structures": (71.0, "medium"),
            "SQL": (68.0, "medium"),
            "Testing": (73.0, "medium"),
        },
    },
    "demo-candidate-003": {
        "job_id": "demo-job-web-001",
        "route": "repository",
        "communication": 74.0,
        "skills": {
            "JavaScript": (79.0, "medium"),
            "React": (66.0, "low"),
            "HTML": (88.0, "high"),
            "Testing": (None, "not_assessed"),
            "Git": (74.0, "medium"),
        },
    },
    "demo-candidate-004": {
        "job_id": "demo-job-python-001",
        "route": "repository",
        "communication": 60.0,
        "skills": {
            "Python": (55.0, "low"),
            "Problem solving": (None, "not_assessed"),
            "Data structures": (None, "not_assessed"),
            "SQL": (61.0, "medium"),
            "Testing": (50.0, "low"),
        },
    },
    "demo-candidate-005": {
        "job_id": "demo-job-web-001",
        "route": "alternative",
        "communication": 76.0,
        "skills": {
            "JavaScript": (71.0, "medium"),
            "React": (74.0, "medium"),
            "HTML": (80.0, "medium"),
            "Testing": (69.0, "medium"),
            "Git": (66.0, "low"),
        },
    },
}


def fixture_proof_cards() -> dict[str, ProofCard]:
    """Build the synthetic Proof Cards shown in the recruiter demo views."""
    cards: dict[str, ProofCard] = {}
    for candidate_id, spec in _FIXTURE_SPECS.items():
        job_id = spec["job_id"]
        job = next(j for j in fx.DEMO_JOBS if j.job_id == job_id)
        requirements = fx.DEMO_JOB_REQUIREMENTS[job_id]
        weights = {r.skill_name: r.weight for r in requirements}
        skill_results = tuple(
            SkillResult(
                skill_name=name,
                performance_score=performance,
                evidence_coverage=100.0 if performance is not None else 0.0,
                evidence_confidence=confidence,
                status="assessed" if performance is not None else "not_assessed",
                evidence=(
                    EvidenceRef(
                        source_type="assessment",
                        skill_name=name,
                        source_id="demo-evidence-001",
                        excerpt="Synthetic demo evidence reference.",
                    ),
                )
                if performance is not None
                else (),
                explanation=(
                    "Synthetic demo result — not a real assessment."
                    if performance is not None
                    else "Not assessed in this synthetic demo."
                ),
                question_count=1 if performance is not None else 0,
            )
            for name, (performance, confidence) in spec["skills"].items()
        )
        performance_score, score_status, coverage = _overall(skill_results, weights)
        cards[candidate_id] = _compose_proof_card(
            candidate_id=candidate_id,
            candidate_reference=fx.DEMO_CANDIDATE_REFERENCES[candidate_id],
            job_id=job_id,
            job_title=job.title,
            route=spec["route"],
            skill_results=skill_results,
            weights=weights,
            performance=performance_score,
            score_status=score_status,
            coverage=coverage,
            evaluations=[],
            profile_skills=[],
            scope=None,
            provider_mode=_PROVIDER_MODE,
            communication_score=spec["communication"],
            provenance={
                "rubric_version": "1.0",
                "route": spec["route"],
                "source": "synthetic fixture",
                "communication_score": spec["communication"],
            },
        )
    return cards
